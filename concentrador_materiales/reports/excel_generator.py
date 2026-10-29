# -*- coding: utf-8 -*-
"""
Generador de archivos Excel para el concentrador de materiales
"""
import os
from datetime import datetime
from typing import Dict, List


from common.logger_config import get_logger
logger = get_logger('concentrador.reports.excel_generator')

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from concentrador_materiales.models.concentrado import ReporteConcentrado, MaterialConcentrado

class ExcelGenerator:
    """
    Genera archivos Excel a partir de un ReporteConcentrado.
    """

    @classmethod
    def generar(cls, reporte: ReporteConcentrado, ruta_archivo: str) -> bool:
        """
        Genera un archivo Excel con el concentrado de materiales.

        Args:
            reporte: ReporteConcentrado con los datos
            ruta_archivo: Ruta donde guardar el archivo

        Returns:
            bool: True si se generó correctamente
        """
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl no está instalado. Ejecute: pip install openpyxl")

        wb = Workbook()

        # Eliminar hoja por defecto
        wb.remove(wb.active)

        # Hoja de resumen
        cls._crear_hoja_resumen(wb, reporte)

        # Hojas por almacén
        for almacen, materiales in reporte.materiales_por_almacen.items():
            if materiales:
                cls._crear_hoja_almacen(wb, almacen, materiales)

        # Guardar archivo
        wb.save(ruta_archivo)
        return True

    @classmethod
    def _crear_hoja_resumen(cls, wb: Workbook, reporte: ReporteConcentrado):
        """Crea la hoja de resumen del reporte."""
        ws = wb.create_sheet("RESUMEN")

        # Estilos
        titulo_font = Font(bold=True, size=14)
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="1f6feb", end_color="1f6feb", fill_type="solid")
        header_font_white = Font(bold=True, color="ffffff")

        # Título
        ws["A1"] = "CONCENTRADOR DE MATERIALES - RESUMEN EJECUTIVO"
        ws["A1"].font = titulo_font
        ws.merge_cells("A1:D1")

        # Fecha
        ws["A3"] = "Fecha de generación: "
        ws["B3"] = reporte.fecha_generacion.strftime("%d/%m/%Y %H:%M:%S")

        # Totales generales
        ws["A5"] = "TOTALES GENERALES"
        ws["A5"].font = header_font

        ws["A6"] = "Presupuestos procesados: "
        ws["B6"] = reporte.total_presupuestos
        ws["A7"] = "Partidas procesadas: "
        ws["B7"] = reporte.total_partidas
        ws["A8"] = "Materiales concentrados: "
        ws["B8"] = reporte.total_materiales

        # Tabla por almacén
        ws["A10"] = "RESUMEN POR ALMACÉN"
        ws["A10"].font = header_font

        headers = ["Almacén", "Materiales", "Unidad", "Cantidad Total"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=11, column=col, value=header)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        row = 12
        for almacen, materiales in reporte.materiales_por_almacen.items():
            if not materiales:
                continue

            # Calcular cantidad total para este almacén
            cantidad_total = sum(m.cantidad for m in materiales)
            unidad = materiales[0].unidad if materiales else ""

            ws.cell(row=row, column=1, value=almacen)
            ws.cell(row=row, column=2, value=len(materiales))
            ws.cell(row=row, column=3, value=unidad)
            ws.cell(row=row, column=4, value=round(cantidad_total, 2))
            row += 1

        # Advertencias
        if reporte.advertencias:
            ws["A" + str(row + 2)] = "ADVERTENCIAS"
            ws["A" + str(row + 2)].font = header_font

            headers_warn = ["Código", "Mensaje"]
            for col, header in enumerate(headers_warn, 1):
                cell = ws.cell(row=row + 3, column=col, value=header)
                cell.font = header_font_white
                cell.fill = header_fill

            warn_row = row + 4
            for adv in reporte.advertencias:
                ws.cell(row=warn_row, column=1, value=adv.get("codigo", ""))
                ws.cell(row=warn_row, column=2, value=adv.get("mensaje", "")[:100])
                warn_row += 1

        # Ajustar anchos
        for col in range(1, 5):
            ws.column_dimensions[get_column_letter(col)].width = 25

    @classmethod
    def _crear_hoja_almacen(cls, wb: Workbook, almacen: str, materiales: List[MaterialConcentrado]):
        """Crea una hoja para un almacén específico."""
        nombre_hoja = almacen[:31]  # Excel máximo 31 caracteres
        ws = wb.create_sheet(nombre_hoja)

        # Estilos
        header_font = Font(bold=True, color="ffffff")
        header_fill = PatternFill(start_color="1f6feb", end_color="1f6feb", fill_type="solid")

        # CORREGIDO: Eliminadas columnas "Dimensiones" y "Origen" según solicitud
        if almacen == "ALUMINIO":
            headers = ["Clave", "Descripción", "Metros Totales (m)", "Largo Tramo (m)", "TRAMOS"]
            campos = ["clave", "descripcion", "metros_totales", "largo_tramo", "tramos_calculados"]
            anchos = [15, 40, 18, 15, 12]
        elif almacen == "VIDRIO":
            headers = ["Clave", "Descripción", "m² Totales", "HOJAS"]
            campos = ["clave", "descripcion", "m2_totales", "hojas_calculadas"]
            anchos = [15, 40, 15, 12]
        elif almacen == "HERRERIA":
            headers = ["Clave", "Descripción", "Metros Totales (m)", "TRAMOS 6m"]
            campos = ["clave", "descripcion", "metros_totales", "tramos_calculados"]
            anchos = [15, 40, 18, 12]
        elif almacen == "HERRAJES":
            headers = ["Clave", "Descripción", "CANTIDAD"]
            campos = ["clave", "descripcion", "cantidad"]
            anchos = [15, 40, 12]
        else:
            # PLASTICOS y OTROS
            headers = ["Clave", "Descripción", "CANTIDAD"]
            campos = ["clave", "descripcion", "cantidad"]
            anchos = [15, 50, 12]

        # Encabezados
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(col)].width = anchos[col - 1]

        # Datos
        for row, material in enumerate(materiales, 2):
            for col, campo in enumerate(campos, 1):
                valor = cls._obtener_valor(material, campo)
                if isinstance(valor, float):
                    valor = round(valor, 2)
                ws.cell(row=row, column=col, value=valor)

        # Totales
        total_row = len(materiales) + 2
        ws.cell(row=total_row, column=1, value="TOTAL")
        ws.cell(row=total_row, column=1).font = Font(bold=True)

        if almacen == "ALUMINIO":
            total_metros = sum(m.metros_totales or 0 for m in materiales)
            total_tramos = sum(m.tramos_calculados or 0 for m in materiales)
            ws.cell(row=total_row, column=3, value=round(total_metros, 2))
            ws.cell(row=total_row, column=5, value=total_tramos)
        elif almacen == "VIDRIO":
            total_m2 = sum(m.m2_totales or 0 for m in materiales)
            total_hojas = sum(m.hojas_calculadas or 0 for m in materiales)
            ws.cell(row=total_row, column=3, value=round(total_m2, 2))
            ws.cell(row=total_row, column=4, value=total_hojas)
        elif almacen == "HERRERIA":
            total_metros = sum(m.metros_totales or 0 for m in materiales)
            total_tramos = sum(m.tramos_calculados or 0 for m in materiales)
            ws.cell(row=total_row, column=3, value=round(total_metros, 2))
            ws.cell(row=total_row, column=4, value=total_tramos)
        else:
            # HERRAJES, PLASTICOS, OTROS
            total_cantidad = sum(m.cantidad for m in materiales)
            ws.cell(row=total_row, column=3, value=round(total_cantidad, 2))

    @classmethod
    def _obtener_valor(cls, material: MaterialConcentrado, campo: str):
        """Obtiene el valor de un campo del material."""
        if campo == "dimensiones" and material.ancho_original and material.alto_original:
            return f"{material.ancho_original:.2f}×{material.alto_original:.2f}"
        if campo == "largo_tramo":
            return material.largo_tramo or 6.10
        valor = getattr(material, campo, None)
        if valor is None and campo == "tramos_calculados":
            return 0
        return valor